import os
import time
import imaplib
import email
import smtplib
import datetime
from pathlib import Path
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from google import genai as g_genai
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from dotenv import load_dotenv

# 🛠️ FORCE ABSOLUTE PATH INJECTION ENVIRONMENT SHIELD
script_dir = Path(__file__).resolve().parent
env_path = script_dir / '.env'
load_dotenv(dotenv_path=env_path, override=True)

# Global App Parameters
OUTPUT_EXCEL_FILENAME = "NIFTY_Financials_Robust_Output.xlsx"
OUTPUT_REPORT_FILENAME = "NIFTY_Portfolio_Screening_Report.docx"

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "")
SENDER_EMAIL = os.environ.get("SENDER_EMAIL", "")
SENDER_PASSWORD = os.environ.get("SENDER_PASSWORD", "")

# Server Connections
SMTP_SERVER = "smtp.gmail.com"
SMTP_PORT = 587
IMAP_SERVER = "imap.gmail.com"
IMAP_PORT = 993

# Initialize Client Instance
client = g_genai.Client(api_key=GEMINI_API_KEY)


def verify_gemini_api_key(client_obj):
    """Performs a lightweight initialization check at boot time to verify key validation."""
    print("🔒 Validating Gemini API key credentials with server infrastructure...")
    try:
        client_obj.models.generate_content(
            model='gemini-2.5-flash',
            contents='Connection probe test.'
        )
        print("✅ Gemini API Key Verification Successful! Gateway connection online.")
        return True
    except Exception as e:
        print("\n❌ CRITICAL ERROR: Gemini API Key Authentication Check Failed!")
        print(f"Error Diagnostic Signature: {e}\n")
        return False


def fetch_latest_reply():
    """Connects via IMAP to find the most recent unread reply referencing the report."""
    try:
        mail = imaplib.IMAP4_SSL(IMAP_SERVER, IMAP_PORT)
        mail.login(SENDER_EMAIL, SENDER_PASSWORD)
        mail.select("inbox")
        
        # Search for UNREAD emails from anyone replying about NIFTY
        status, messages = mail.search(None, '(UNSEEN SUBJECT "NIFTY")')
        if status != "OK" or not messages[0]:
            mail.logout()
            return None
            
        latest_message_id = messages[0].split()[-1]
        status, data = mail.fetch(latest_message_id, "(RFC822)")
        
        raw_email = data[0][1]
        msg = email.message_from_bytes(raw_email)
        
        sender = msg.get("From")
        subject = msg.get("Subject")
        
        email_body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    email_body = part.get_payload(decode=True).decode()
                    break
        else:
            email_body = msg.get_payload(decode=True).decode()
            
        print(f"📥 New Reply Received from {sender}: '{subject}'")
        return {"sender": sender, "body": email_body, "id": latest_message_id, "mail_obj": mail}
    except Exception as e:
        print(f"⚠️ IMAP Folder Access Alert: {e}")
        return None


def apply_ai_modifications(feedback_text, max_retries=3, initial_delay=5):
    """Feeds user feedback to Gemini to rebuild the analysis with a retry shield."""
    print("🧠 Orchestrating Gemini to process modifications based on reply...")
    
    original_text = ""
    if os.path.exists(OUTPUT_REPORT_FILENAME):
        try:
            doc_obj = Document(OUTPUT_REPORT_FILENAME)
            original_text = "\n".join([p.text for p in doc_obj.paragraphs])
        except:
            pass

    modification_prompt = f"""
    You are the Senior Forensic Risk Auditor. A user reviewed your report and sent this feedback:
    "{feedback_text}"
    
    Here is the context text of the existing report:
    {original_text}
    
    Task: Regenerate the RATIONALE and RATING section for the specific company mentioned in the feedback. 
    Incorporate their requested adjustments intelligently, using financial reasoning.
    
    Output Format STRICTLY as followed:
    TARGET_COMPANY: [Exact Name of Company to Change]
    RATING: [NEW BUY, SELL, or HOLD]
    REASONING: [New dense, high-impact 4-5 sentence justification paragraph citing metrics]
    """
    
    delay = initial_delay
    for attempt in range(max_retries):
        try:
            response = client.models.generate_content(
                model='gemini-2.5-flash',
                contents=modification_prompt
            )
            return response.text
        except Exception as e:
            error_str = str(e)
            if "503" in error_str or "UNAVAILABLE" in error_str.upper():
                print(f"⏳ [Attempt {attempt + 1}/{max_retries}] Heavy server load (503). Retrying in {delay}s...")
                time.sleep(delay)
                delay *= 2
            else:
                print(f"❌ Gemini optimization runtime error: {e}")
                return None
    return None


def update_word_document(ai_update_text):
    """
    Appends a brand-new, beautifully formatted 'Amended Review Section' 
    directly to the bottom of the existing Word Document to guarantee 
    modifications take effect visibly without search mismatches.
    """
    if not ai_update_text or "TARGET_COMPANY:" not in ai_update_text:
        print("⚠️ Invalid AI text format layout. Skipping Word alteration ledger.")
        return False
        
    try:
        # Extract individual lines cleanly
        lines = ai_update_text.strip().split('\n')
        target_company = "Unknown Asset"
        new_rating = "AMENDED"
        new_reasoning = ""
        
        for line in lines:
            if "TARGET_COMPANY:" in line.upper():
                target_company = line.split(":", 1)[1].strip()
            elif "RATING:" in line.upper():
                new_rating = line.split(":", 1)[1].strip()
            elif "REASONING:" in line.upper():
                new_reasoning = line.split(":", 1)[1].strip()

        # Fallback multi-line lookup if reasoning got chunked
        if not new_reasoning:
            idx = ai_update_text.upper().find("REASONING:")
            if idx != -1:
                new_reasoning = ai_update_text[idx + len("REASONING:"):].strip()

        print(f"💾 Opening report to append amendment ledger section for {target_company}...")
        
        # Load up the master document file
        doc = Document(OUTPUT_REPORT_FILENAME)
        
        # 1. Inject a clean visual section line divider
        doc.add_paragraph("\n" + "_"*60 + "\n")
        
        # 2. Add an amendment ledger section header
        heading = doc.add_heading("🔄 Dynamic AI Portfolio Amendment Ledger", level=2)
        heading.runs[0].font.color.rgb = RGBColor(0, 51, 102)  # Deep Executive Navy Color
        
        # 3. Compile asset metadata details
        p_meta = doc.add_paragraph()
        p_meta.add_run("📍 Target Security: ").bold = True
        p_meta.add_run(f"{target_company}\n")
        p_meta.add_run("📊 Updated Committee Action: ").bold = True
        
        # Color code the recommendation text natively
        run_rating = p_meta.add_run(f"[{new_rating.upper()}]")
        run_rating.bold = True
        if "BUY" in new_rating.upper():
            run_rating.font.color.rgb = RGBColor(0, 128, 0)     # Green
        elif "SELL" in new_rating.upper():
            run_rating.font.color.rgb = RGBColor(204, 0, 0)    # Red
        else:
            run_rating.font.color.rgb = RGBColor(204, 153, 0)   # Amber / Gold for Hold
            
        timestamp_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        p_meta.add_run(f"\n📅 Audit Lifecycle Timestamp: {timestamp_str} (Processed via Mail Automation Loop)\n")
        
        # 4. Inject structural Blockquote style for the fresh rationale text
        p_title = doc.add_paragraph()
        p_title.add_run("🧠 REVISED FORENSIC AUDITOR RATIONALE:").bold = True
        
        p_reason = doc.add_paragraph()
        p_reason.paragraph_format.left_indent = Inches(0.5)  # Indent text like an executive blockquote
        run_reason = p_reason.add_run(f'"{new_reasoning}"')
        run_reason.italic = True
        run_reason.font.size = Pt(10.5)
        run_reason.font.color.rgb = RGBColor(64, 64, 64)      # Dark Charcoal Gray
        
        # Save back onto the file path
        doc.save(OUTPUT_REPORT_FILENAME)
        print("✅ Success! Section successfully appended and saved directly to the bottom of the Word report.")
        return True
    except Exception as e:
        print(f"⚠️ Failed to write amendment modifications into Word file: {e}")
        return False


def send_updated_response(recipient_email):
    """Emails the freshly modified report back to the reviewer."""
    print(f"📤 Dispatching updated report attachment directly to {recipient_email}...")
    try:
        msg = MIMEMultipart()
        msg['From'] = SENDER_EMAIL
        msg['To'] = recipient_email
        msg['Subject'] = "Re: [AMENDED REPORT] NIFTY Portfolio Executive Screening Document"

        body = (
            "Hello,\n\nOur Agentic Pipeline has processed your response metrics, "
            "re-evaluated the structural positions via Gemini, and appended an official "
            "Amended Ledger Section directly to the bottom of your report document.\n\n"
            "Please find the revised .docx master file attached for your review.\n\n"
            "Best Regards,\nQuantAgent Pipeline Framework"
        )
        msg.attach(MIMEText(body, 'plain'))

        with open(OUTPUT_REPORT_FILENAME, "rb") as attachment:
            part = MIMEBase('application', 'octet-stream')
            part.set_payload(attachment.read())
            encoders.encode_base64(part)
            part.add_header('Content-Disposition', f"attachment; filename={os.path.basename(OUTPUT_REPORT_FILENAME)}")
            msg.attach(part)

        server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
        server.starttls()
        server.login(SENDER_EMAIL, SENDER_PASSWORD)
        server.sendmail(SENDER_EMAIL, recipient_email, msg.as_string())
        server.quit()
        print("✅ Reply successfully transmitted across network channels.")
    except Exception as e:
        print(f"⚠️ SMTP Protocol Transmission Error: {e}")


# =====================================================================
# SYSTEM MAIN ENGINE RUNTIME POLLING LOOP
# =====================================================================
if __name__ == "__main__":
    print("🚀 Initializing QuantAgent Active Background Services...")
    
    if not verify_gemini_api_key(client):
        print("🛑 Startup Aborted: Pipeline halted due to invalid API key configurations.")
        exit(1)
        
    print("📬 Continuous Listener Loop active and scanning for replies... (Press Ctrl+C to halt)")
    while True:
        reply_data = fetch_latest_reply()
        
        if reply_data:
            update_instructions = apply_ai_modifications(reply_data["body"])
            success = update_word_document(update_instructions)
            
            if success:
                send_updated_response(reply_data["sender"])
                
                # Mark email as read so we don't process it on the next loop
                try:
                    reply_data["mail_obj"].store(reply_data["id"], '+FLAGS', '\\Seen')
                    reply_data["mail_obj"].logout()
                except:
                    pass
        else:
            time.sleep(30)