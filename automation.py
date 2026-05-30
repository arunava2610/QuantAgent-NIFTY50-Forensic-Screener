import os
import time
import argparse
import schedule
import pandas as pd
import yfinance as yf
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

# =====================================================================
# CONFIGURATION & CONSTANTS
# =====================================================================
INPUT_TICKERS_FILENAME = "tickers.xlsx"
OUTPUT_EXCEL_FILENAME = "NIFTY_Financials_Robust_Output.xlsx"
TARGET_YEARS = [2021, 2022, 2023, 2024, 2025, 2026]

# Standard Accounting Rows (From your main.py)
balance_sheet_items = [
    "Cash And Cash Equivalents", "Short Term Investments", "Net Receivables", "Inventory",
    "Other Current Assets", "Total Current Assets", "Property Plant And Equipment", "Goodwill",
    "Total Assets", "Accounts Payable", "Short Term Debt", "Other Current Liabilities",
    "Total Current Liabilities", "Long Term Debt", "Deferred Tax Liabilities",
    "Other Non Current Liabilities", "Total Non Current Liabilities", "Total Liabilities",
    "Ordinary Shares Number", "Retained Earnings", "Total Stockholder Equity",
    "Total Liabilities And Stockholders Equity"
]

income_statement_items = [
    "Total Revenue", "Cost Of Revenue", "Gross Profit", "Operating Expense",
    "Operating Income", "Net Non Operating Interest Income Expense", "Other Income Expense",
    "Pretax Income", "Income Tax Expense", "Net Income", "Normalized Income",
    "Interest Income", "Interest Expense", "Ebit", "Ebitda"
]

cash_flow_items = [
    "Operating Cash Flow", "Investing Cash Flow", "Financing Cash Flow",
    "Net Cash Flow", "Depreciation And Amortization", "Stock Based Compensation",
    "Change In Working Capital", "Capital Expenditure", "Issuance Of Stock",
    "Repurchase Of Stock", "Issuance Of Debt", "Repayment Of Debt"
]

fiscal_quarters = {
    2021: [(6, 2020), (9, 2020), (12, 2020), (3, 2021)],
    2022: [(6, 2021), (9, 2021), (12, 2021), (3, 2022)],
    2023: [(6, 2022), (9, 2022), (12, 2022), (3, 2023)],
    2024: [(6, 2023), (9, 2023), (12, 2023), (3, 2024)],
    2025: [(6, 2024), (9, 2024), (12, 2024), (3, 2025)],
    2026: [(6, 2025), (9, 2025), (12, 2025), (3, 2026)]
}

# =====================================================================
# EXTRACTION HELPERS
# =====================================================================
def safe_extract_historical(df_source, item_name, target_col):
    """Safely extracts historical data by matching specific date columns."""
    if df_source is None or df_source.empty:
        return None
    if item_name in df_source.index and target_col in df_source.columns:
        return df_source.loc[item_name, target_col]
    
    lower_index = [str(x).lower().strip() for x in df_source.index]
    target_lower = item_name.lower().strip()
    if target_lower in lower_index:
        real_idx = df_source.index[lower_index.index(target_lower)]
        if target_col in df_source.columns:
            return df_source.loc[real_idx, target_col]
    return None

def safe_extract_latest(df_source, item_name):
    """Safely extracts the most recent data point for the Ratios sheet."""
    if df_source is None or df_source.empty:
        return 0.0
    try:
        match = df_source[df_source.index.str.contains(item_name, case=False, na=False)]
        if not match.empty:
            val = match.iloc[0, 0]
            return float(val) if pd.notna(val) else 0.0
    except:
        pass
    return 0.0


# =====================================================================
# SHEET FORMATTING WRITER
# =====================================================================
def write_historical_sheet(wb, year, data):
    """Formats and writes the historical year-by-year data into beautiful Excel sheets."""
    if not data:
        return
    ws = wb.create_sheet(str(year))
    df_sheet = pd.DataFrame(data)
    df_sheet = df_sheet.sort_values(by=["Company Name", "Quarter Sort"]).drop(columns=["Quarter Sort"])

    for r_idx, row in enumerate(dataframe_to_rows(df_sheet, index=False, header=True), 3):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    num_meta_cols = 4  
    start_pl = num_meta_cols + 1
    start_bs = start_pl + len(income_statement_items)
    start_cf = start_bs + len(balance_sheet_items)

    ws.merge_cells(start_row=1, start_column=start_pl, end_row=1, end_column=start_bs - 1)
    ws.merge_cells(start_row=1, start_column=start_bs, end_row=1, end_column=start_cf - 1)
    ws.merge_cells(start_row=1, start_column=start_cf, end_row=1, end_column=start_cf + len(cash_flow_items) - 1)

    ws.cell(row=1, column=start_pl).value = "Income Statement (P&L)"
    ws.cell(row=1, column=start_bs).value = "Balance Sheet"
    ws.cell(row=1, column=start_cf).value = "Cash Flow"

    fill = PatternFill(start_color="B7D7F7", end_color="B7D7F7", fill_type="solid")
    font = Font(bold=True, size=11)
    for col in [start_pl, start_bs, start_cf]:
        cell = ws.cell(row=1, column=col)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = fill
        cell.font = font

    end_col_letter = get_column_letter(df_sheet.shape[1])
    ws.auto_filter.ref = f"A3:{end_col_letter}3"
    
    for i in range(1, df_sheet.shape[1] + 1):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = 22
        if i == 3:
            for row in range(4, ws.max_row + 1):
                ws[f"{col_letter}{row}"].number_format = "YYYY-MM-DD"
        elif i > 4:
            for row in range(4, ws.max_row + 1):
                cell = ws[f"{col_letter}{row}"]
                if cell.value is not None:
                    cell.number_format = '#,##0.00'

# =====================================================================
# CORE SCRAPING ENGINE
# =====================================================================
def generate_financial_excel():
    print(f"\n📊 [{time.strftime('%Y-%m-%d %H:%M:%S')}] Launching Full Historical & Ratio Data Engine...")
    
    if not os.path.exists(INPUT_TICKERS_FILENAME):
        print(f"❌ ERROR: '{INPUT_TICKERS_FILENAME}' not found. Place it in the root folder!")
        return

    df_tickers = pd.read_excel(INPUT_TICKERS_FILENAME)
    ticker_col = 'Yahoo_Ticker' if 'Yahoo_Ticker' in df_tickers.columns else ('Symbol' if 'Symbol' in df_tickers.columns else df_tickers.columns[0])
    company_col = 'Company Name' if 'Company Name' in df_tickers.columns else df_tickers.columns[0]

    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
        
    historical_results = {year: [] for year in TARGET_YEARS}
    ratio_data = []

    for idx, row in df_tickers.iterrows():
        ticker = str(row[ticker_col]).strip()
        company = row[company_col]
        if not ticker.endswith('.NS') and not ticker.endswith('.BO'):
            ticker += '.NS'
            
        print(f"[{idx+1}/{len(df_tickers)}] Processing Comprehensive Matrix for: {company}...")
        
        try:
            stock = yf.Ticker(ticker)
            
            # Pull all dataframes at once to speed up processing
            ann_fin, ann_bs, ann_cf = stock.financials, stock.balance_sheet, stock.cashflow
            q_fin, q_bs, q_cf = stock.quarterly_financials, stock.quarterly_balance_sheet, stock.quarterly_cashflow

            # ---------------------------------------------------------
            # 1. BUILD THE LATEST RATIOS
            # ---------------------------------------------------------
            net_inc = safe_extract_latest(ann_fin, "Net Income")
            rev = safe_extract_latest(ann_fin, "Total Revenue")
            gross = safe_extract_latest(ann_fin, "Gross Profit")
            assets = safe_extract_latest(ann_bs, "Total Assets")
            equity = safe_extract_latest(ann_bs, "Total Stockholder Equity")
            debt = safe_extract_latest(ann_bs, "Long Term Debt")
            curr_assets = safe_extract_latest(ann_bs, "Total Current Assets")
            curr_liab = safe_extract_latest(ann_bs, "Total Current Liabilities")
            cash = safe_extract_latest(ann_bs, "Cash And Cash Equivalents")
            ocf = safe_extract_latest(ann_cf, "Operating Cash Flow")

            npm = (net_inc / rev * 100) if rev else 0
            gm = (gross / rev * 100) if rev else 0
            roa = (net_inc / assets * 100) if assets else 0
            roe = (net_inc / equity * 100) if equity else 0
            dte = (debt / equity) if equity else 0
            dta = (debt / assets * 100) if assets else 0
            cr = (curr_assets / curr_liab) if curr_liab else 0
            ocf_debt = (ocf / debt) if debt else 0
            cta = (cash / assets * 100) if assets else 0

            ratio_data.append([
                company, ticker, 
                round(npm, 2), round(gm, 2), round(roa, 2), round(roe, 2),
                round(dte, 2), round(dta, 2), round(cr, 2), round(ocf_debt, 2), round(cta, 2)
            ])

            # ---------------------------------------------------------
            # 2. BUILD HISTORICAL SHEETS (2021 - 2026)
            # ---------------------------------------------------------
            if ann_fin is not None and not ann_fin.empty:
                for fy in TARGET_YEARS:
                    ann_cols = [col for col in ann_fin.columns if hasattr(col, "year") and col.year == fy]
                    if not ann_cols:
                        continue
                    ann_col = ann_cols[0]
                    
                    ann_entry = {
                        'Company Name': company, 'Ticker': ticker,
                        'Last Updated': str(ann_col.date()), 'Type': "Annual", 'Quarter Sort': 0
                    }
                    for item in income_statement_items:
                        ann_entry[item] = safe_extract_historical(ann_fin, item, ann_col)
                    for item in balance_sheet_items:
                        ann_entry[item] = safe_extract_historical(ann_bs, item, ann_col)
                    for item in cash_flow_items:
                        ann_entry[item] = safe_extract_historical(ann_cf, item, ann_col)
                    historical_results[fy].append(ann_entry)

            if q_fin is not None and not q_fin.empty:
                q_fin_cols = pd.to_datetime(q_fin.columns)
                for fy, quarter_list in fiscal_quarters.items():
                    for q_idx, (target_month, target_year) in enumerate(quarter_list):
                        matched_cols = [col for col in q_fin_cols if col.month == target_month and col.year == target_year]
                        if not matched_cols:
                            continue
                        real_quarter_date = matched_cols[0]
                        orig_col_key = q_fin.columns[q_fin_cols.get_loc(real_quarter_date)]
                        
                        entry = {
                            'Company Name': company, 'Ticker': ticker,
                            'Last Updated': str(real_quarter_date.date()), 'Type': f"Q{q_idx+1}", 'Quarter Sort': q_idx + 1
                        }
                        for item in income_statement_items:
                            entry[item] = safe_extract_historical(q_fin, item, orig_col_key)
                        for item in balance_sheet_items:
                            entry[item] = safe_extract_historical(q_bs, item, orig_col_key)
                        for item in cash_flow_items:
                            entry[item] = safe_extract_historical(q_cf, item, orig_col_key)
                        historical_results[fy].append(entry)

            time.sleep(1.2) # Throttle to avoid Yahoo Finance IP blocks

        except Exception as e:
            print(f"   ⚠️ Warning on {ticker}: {e}")
            time.sleep(2)

    # ---------------------------------------------------------
    # 3. WRITE ALL DATA TO EXCEL WORKBOOK
    # ---------------------------------------------------------
    # Write Ratios Tab First
    ws_ratios = wb.create_sheet("Ratios", 0) 
    ratio_headers = [
        "Company Name", "Ticker", "Net Profit Margin (%)", "Gross Margin (%)", 
        "Return on Assets (%)", "Return on Equity (%)", "Debt to Equity", 
        "Debt to Assets (%)", "Current Ratio", "Op Cash Flow to Debt", "Cash to Assets (%)"
    ]
    ws_ratios.append(ratio_headers)
    
    fill = PatternFill(start_color="B7D7F7", end_color="B7D7F7", fill_type="solid")
    for col in range(1, len(ratio_headers) + 1):
        cell = ws_ratios.cell(row=1, column=col)
        cell.fill = fill
        cell.font = Font(bold=True)
        ws_ratios.column_dimensions[get_column_letter(col)].width = 20

    for r_data in ratio_data:
        ws_ratios.append(r_data)
        
    for row in range(2, ws_ratios.max_row + 1):
        for col in range(3, 12):
            ws_ratios.cell(row=row, column=col).number_format = '#,##0.00'

    # Write Historical Tabs (2021-2026)
    for y in TARGET_YEARS:
        full_data = [row for row in historical_results[y] if row['Type'] == "Annual" or row['Type'] in ["Q1", "Q2", "Q3", "Q4"]]
        write_historical_sheet(wb, y, full_data)

    wb.save(OUTPUT_EXCEL_FILENAME)
    print(f"\n✅ Heavy Scraping Complete. Full matrix + Ratios compiled successfully into {OUTPUT_EXCEL_FILENAME}.")


# =====================================================================
# AUTOMATION DAEMON SCHEDULER
# =====================================================================
def daily_automation_job():
    generate_financial_excel()
    print(f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] Hibernating until tomorrow 12:01 AM.\n")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--manual", action="store_true", help="Force run immediately")
    args = parser.parse_args()

    if args.manual:
        print("⚡ Manual Override Triggered.")
        daily_automation_job()
    else:
        print("⏳ Automation Daemon Started. Scheduled for 12:01 AM daily.")
        schedule.every().day.at("00:01").do(daily_automation_job)
        while True:
            schedule.run_pending()
            time.sleep(60)