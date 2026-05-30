from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
import docx
import os

# Initialize the FastAPI app
app = FastAPI(title="QuantAgent API")

# Configure CORS so your Vercel React app can communicate with this backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # In production, replace "*" with your Vercel URL for tighter security
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Define the data structure expected from the React frontend POST request
class ReportRequest(BaseModel):
    companies: list[str]
    custom_feedback: str

# ---------------------------------------------------------
# ENDPOINT 1: Feed the UI with Nifty 50 Companies
# ---------------------------------------------------------
@app.get("/api/companies")
def get_companies():
    # A starter list of Nifty 50 equities to populate your React grid
    return [
        {"ticker": "RELIANCE.NS", "name": "Reliance Industries"},
        {"ticker": "TCS.NS", "name": "Tata Consultancy Services"},
        {"ticker": "HDFCBANK.NS", "name": "HDFC Bank"},
        {"ticker": "INFY.NS", "name": "Infosys"},
        {"ticker": "ICICIBANK.NS", "name": "ICICI Bank"},
        {"ticker": "ITC.NS", "name": "ITC Limited"}
    ]

# ---------------------------------------------------------
# ENDPOINT 2: Generate and Download the Word Report
# ---------------------------------------------------------
@app.post("/api/generate-report")
def generate_report(request: ReportRequest):
    # 1. Initialize a blank MS Word document
    doc = docx.Document()
    
    # 2. Add content to the document based on UI inputs
    doc.add_heading('QuantAgent Portfolio Rationale', 0)
    
    doc.add_heading('Target Securities Evaluated:', level=1)
    for ticker in request.companies:
        doc.add_paragraph(f"• {ticker}", style='List Bullet')
        
    doc.add_heading('Strategic Guidelines Applied:', level=1)
    doc.add_paragraph(request.custom_feedback if request.custom_feedback else "Standard screening matrices applied.")
    
    # Placeholder for where your actual AI generation logic will go
    doc.add_heading('AI Analysis Output', level=1)
    doc.add_paragraph("This is a placeholder for the Gemini AI output. The backend successfully received your configuration and compiled this document.")

    # 3. Save the document temporarily to the server
    file_path = "temp_report.docx"
    doc.save(file_path)
    
    # 4. Stream the file back to the React UI as a downloadable blob
    return FileResponse(
        path=file_path, 
        filename="QuantAgent_Custom_Report.docx", 
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )