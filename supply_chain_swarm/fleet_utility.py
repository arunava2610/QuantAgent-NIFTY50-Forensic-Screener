import pandas as pd
import requests
import sys
import os
import io
import base64
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import logger, ERP_DRIVE_URL

def run_fleet_assignment(allocated_node: str):
    """Assigns an idle fleet asset, highlights the row, and returns a Base64 in-memory file."""
    try:
        # 1. Fetch live database from Drive into memory
        headers = {'User-Agent': 'Mozilla/5.0'}
        res = requests.get(ERP_DRIVE_URL, headers=headers)
        if res.status_code != 200:
            raise Exception(f"Drive connection failed: {res.status_code}")

        file_stream = io.BytesIO(res.content)

        # 2. Open workbook in memory
        wb = load_workbook(file_stream)
        if "Active_Fleet" not in wb.sheetnames:
            raise Exception("Active_Fleet sheet missing.")
            
        ws = wb["Active_Fleet"]
        green_fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")

        # 3. Find IDLE asset
        target_row = None
        assigned_asset_id = None
        
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=5).value == "IDLE":
                target_row = row
                assigned_asset_id = ws.cell(row=row, column=1).value
                break
                
        if not target_row:
            return "No idle fleet assets available.", None

        # 4. Update data & highlight row
        ws.cell(row=target_row, column=4).value = allocated_node
        ws.cell(row=target_row, column=5).value = "EN_ROUTE"
        ws.cell(row=target_row, column=6).value = "100%"
        
        for col in range(1, 9):
            ws.cell(row=target_row, column=col).fill = green_fill

        # 5. Save modified workbook back to a NEW in-memory byte stream
        output_stream = io.BytesIO()
        wb.save(output_stream)
        output_stream.seek(0) # Reset stream position to the beginning

        # 6. Encode to Base64 so it can be sent via JSON to React
        b64_string = base64.b64encode(output_stream.read()).decode('utf-8')

        return assigned_asset_id, b64_string

    except Exception as e:
        logger.error(f"Fleet Assignment Error: {str(e)}")
        raise e