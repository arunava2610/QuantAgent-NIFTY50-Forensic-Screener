import os
import time
import json
import requests
import smtplib
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import docx
from docx.shared import Pt, RGBColor, Inches
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from fastapi import APIRouter
from fastapi.responses import FileResponse
from pydantic import BaseModel
from google import genai as g_genai

from config import logger, EXCEL_DRIVE_URL, get_api_key

router = APIRouter()

class ReportRequest(BaseModel):
    companies: list[str]
    custom_feedback: str
    email: str = "" 

cached_companies = []
master_financial_df = None

def download_excel_from_drive():
    temp_file_path = "temp_drive_matrix.xlsx"
    try:
        headers = {'User-Agent': 'Mozilla/5.0'}
        response = requests.get(EXCEL_DRIVE_URL, headers=headers)
        if response.status_code == 200:
            with open(temp_file_path, "wb") as f:
                f.write(response.content)
            return temp_file_path
    except Exception as e:
        logger.error(f"Failed to download Excel from Drive: {e}")
    return None

def send_report_via_email(recipient_email, file_path):
    sender_email = os.environ.get("SENDER_EMAIL")
    sender_password = os.environ.get("SENDER_PASSWORD")
    if not sender_email or not sender_password:
        raise Exception("SMTP credentials not configured.")

    try:
        msg = MIMEMultipart()
        msg['From'] = sender_email
        msg['To'] = recipient_email
        msg['Subject'] = "NIFTY Investment Committee Executive Screening Document"
        body = "Please find attached the finalized executive screening report."
        msg.attach(MIMEText(body, 'plain'))

        with open(file_path, "rb") as attachment:
            part = MIMEBase('application', 'octet-stream')
            part.set_payload(attachment.read())
            encoders.encode_base64(part)
            part.add_header('Content-Disposition', f"attachment; filename={os.path.basename(file_path)}")
            msg.attach(part)

        server = smtplib.SMTP("smtp.gmail.com", 587)
        server.starttls()
        server.login(sender_email, sender_password)
        server.sendmail(sender_email, recipient_email, msg.as_string())
        server.quit()
    except Exception as e:
        logger.error(f"Failed to send email: {e}")
        raise e

@router.get("/api/companies")
def get_companies():
    global cached_companies
    if cached_companies:
        return cached_companies
    try:
        local_excel_path = download_excel_from_drive()
        df = pd.read_excel(local_excel_path, sheet_name="Ratios") # type: ignore
        cached_companies = [{"ticker": str(row["Ticker"]), "name": str(row["Company Name"])} for idx, row in df.iterrows()]
        return cached_companies
    except Exception as e:
        return [{"ticker": "RELIANCE.NS", "name": "Reliance Industries"}]

@router.post("/api/generate-report")
def generate_report(request: ReportRequest):
    global master_financial_df
    live_api_key = get_api_key()
    if not live_api_key:
        return {"error": "Failed to retrieve Gemini API Key."}
    
    client = g_genai.Client(api_key=live_api_key) # type: ignore

    if master_financial_df is None:
        local_excel_path = download_excel_from_drive()
        xls = pd.ExcelFile(local_excel_path) # type: ignore
        all_years_data = []
        for sheet_name in xls.sheet_names:
            if sheet_name != "Ratios":
                year_df = pd.read_excel(xls, sheet_name=sheet_name, header=2) # type: ignore
                year_df['Fiscal_Year'] = int(sheet_name)
                all_years_data.append(year_df)
        master_financial_df = pd.concat(all_years_data, ignore_index=True)

    compiled_prompt = "You are an Elite Corporate Investment Committee Member. Evaluate the following securities.\n\n"
    company_metadata = [] 

    for ticker in request.companies:
        comp_annual_df = master_financial_df[(master_financial_df['Ticker'] == ticker) & (master_financial_df['Type'] == 'Annual')].copy()
        if comp_annual_df.empty: continue
            
        company_name = comp_annual_df['Company Name'].iloc[0]
        comp_annual_df = comp_annual_df.sort_values('Fiscal_Year')
        comp_annual_df['Net Profit Margin (%)'] = (comp_annual_df['Net Income'] / comp_annual_df['Total Revenue']) * 100
        comp_annual_df['Return on Assets (%)'] = (comp_annual_df['Net Income'] / comp_annual_df['Total Assets']) * 100
        comp_annual_df['Debt to Assets (%)'] = (comp_annual_df['Long Term Debt'].fillna(0) / comp_annual_df['Total Assets']) * 100
        
        ratios_summary_text = ""
        valid_ratio_df = comp_annual_df.dropna(subset=['Net Profit Margin (%)', 'Total Revenue'])
        for _, r_row in valid_ratio_df.iterrows():
            ratios_summary_text += f"• FY {int(r_row['Fiscal_Year'])} | NPM: {r_row['Net Profit Margin (%)']:.2f}% | ROA: {r_row['Return on Assets (%)']:.2f}% | Debt/Assets: {r_row['Debt to Assets (%)']:.2f}%\n"

        comp_all_df = master_financial_df[master_financial_df['Ticker'] == ticker].dropna(how='all', axis=1)
        compressed_ledger_csv = comp_all_df.to_csv(index=False)

        company_metadata.append({"ticker": ticker, "name": company_name, "valid_ratio_df": valid_ratio_df})
        compiled_prompt += f"--- TARGET SECURITY: {ticker} ({company_name}) ---\n[HISTORICAL LEDGER]\n{compressed_ledger_csv}\n[RATIOS PROFILE]\n{ratios_summary_text}\n\n"

    compiled_prompt += "Output STRICTLY: TICKER: [X] \n RATING: [BUY/SELL/HOLD] \n REASONING: [Text] \n==="

    max_retries = 3
    ai_response = ""
    for attempt in range(1, max_retries + 1):
        try:
            logger.info(f"Sending payload to Gemini (Attempt {attempt}/{max_retries})...")
            response = client.models.generate_content(model='gemini-2.5-flash', contents=compiled_prompt) # type: ignore
            ai_response = response.text
            logger.info("✅ Gemini API generation successful.")
            break 
        except Exception as e:
            error_msg = str(e)
            logger.warning(f"⚠️ Gemini API failed on attempt {attempt}. Error: {error_msg}")
            
            if "API key" in error_msg or "400" in error_msg or "401" in error_msg or "403" in error_msg:
                logger.error("🚨 CRITICAL: Invalid or missing Gemini API Key detected! Aborting generation.")
                break 
                
            if attempt < max_retries: 
                logger.info("Waiting 5 seconds before retrying...")
                time.sleep(5)
            else:
                logger.error("❌ Max retries reached. AI Generation completely failed.")

    ai_results = {}
    try:
        cleaned_json = ai_response.strip().replace("```json", "").replace("```", "")
        parsed_json = json.loads(cleaned_json)
        if isinstance(parsed_json, list):
            for item in parsed_json:
                raw_ticker = str(item.get("ticker", "")).strip().upper()
                if not raw_ticker.endswith('.NS') and not raw_ticker.endswith('.BO'): raw_ticker += '.NS'
                if raw_ticker: ai_results[raw_ticker] = {"rating": str(item.get("rating", "")).upper(), "reasoning": str(item.get("reasoning", "")).strip()}
    except json.JSONDecodeError:
        current_ticker, current_rating, current_reasoning = None, "REVIEW REQUIRED", ""
        for line in ai_response.split('\n'):
            clean_line = line.strip().replace("*", "").replace("#", "")
            upper_line = clean_line.upper()
            if upper_line.startswith("TICKER:"):
                if current_ticker: ai_results[current_ticker] = {"rating": current_rating, "reasoning": current_reasoning.strip()}
                raw_ticker = clean_line.split(":", 1)[1].strip().upper()
                if not raw_ticker.endswith('.NS') and not raw_ticker.endswith('.BO'): raw_ticker += '.NS'
                current_ticker, current_reasoning, current_rating = raw_ticker, "", "REVIEW REQUIRED"
            elif upper_line.startswith("RATING:"): current_rating = clean_line.split(":", 1)[1].strip()
            elif upper_line.startswith("REASONING:"): current_reasoning = clean_line.split(":", 1)[1].strip() + " "
            elif current_ticker and clean_line and upper_line != "===": current_reasoning += clean_line + " "
        if current_ticker: ai_results[current_ticker] = {"rating": current_rating, "reasoning": current_reasoning.strip()}

    doc = docx.Document() # type: ignore
    doc.add_heading('NIFTY Executive Investment Screening Summary', level=0) # type: ignore

    for idx, meta in enumerate(company_metadata):
        ticker, company_name, valid_ratio_df = meta["ticker"], meta["name"], meta["valid_ratio_df"]
        result = ai_results.get(ticker, {"rating": "HOLD", "reasoning": "AI Generation skipped."})
        
        chart_filename = f"temp_chart_{idx}.png"
        try:
            fig, ax = plt.subplots(figsize=(7, 4))
            years_str = valid_ratio_df['Fiscal_Year'].astype(str)
            ax.plot(years_str, valid_ratio_df['Net Profit Margin (%)'], marker='o', label='NPM %')
            ax.plot(years_str, valid_ratio_df['Return on Assets (%)'], marker='s', label='ROA %')
            ax.plot(years_str, valid_ratio_df['Debt to Assets (%)'], marker='^', label='Debt/Assets %')
            ax.legend(loc='best')
            plt.tight_layout()
            plt.savefig(chart_filename, dpi=150)
            plt.close(fig)
        except Exception: pass

        p_heading = doc.add_paragraph() # type: ignore
        p_heading.add_run(f"\n🏢 {company_name} | Action: ").bold = True # type: ignore
        r_rate = p_heading.add_run(result["rating"]) # type: ignore
        r_rate.bold = True
        
        if "BUY" in result["rating"].upper(): r_rate.font.color.rgb = RGBColor(0, 128, 0) # type: ignore
        elif "SELL" in result["rating"].upper(): r_rate.font.color.rgb = RGBColor(180, 0, 0) # type: ignore

        doc.add_paragraph("🧠 Forensic Rationale:").bold = True # type: ignore
        doc.add_paragraph(result["reasoning"]) # type: ignore

        if os.path.exists(chart_filename):
            doc.add_paragraph().add_run().add_picture(chart_filename, width=Inches(5.5)) # type: ignore
            os.remove(chart_filename) 

    output_path = "NIFTY_Custom_Screening_Report.docx"
    doc.save(output_path) # type: ignore
    
    if request.email and "@" in request.email:
        try:
            send_report_via_email(request.email, output_path)
            os.remove(output_path) 
            return {"message": f"Dispatched to {request.email}"}
        except Exception as e:
            return {"error": f"Failed to send email: {e}"}
    else:
        return FileResponse(path=output_path, filename="NIFTY_Custom_Screening_Report.docx", media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')